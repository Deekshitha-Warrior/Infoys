"""Integration tests use a fresh in-memory DB and synthetic resume fixtures."""

import io
import json
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from app.main import app
from app.database import Base, get_db
from app.intelligence import canonical, predict, resume_analysis
from app.parser import parse_resume


@pytest.fixture
def client(tmp_path, monkeypatch):
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)

    def db_override():
        with Session() as db:
            yield db

    app.dependency_overrides[get_db] = db_override
    monkeypatch.setattr("app.routers.resume.UPLOAD_DIR", tmp_path)
    with TestClient(app) as client:
        yield client
    app.dependency_overrides.clear()
    engine.dispose()


def account(client, email="test@example.com"):
    response = client.post(
        "/api/auth/register",
        json={
            "name": "Test Student",
            "email": email,
            "password": "testing123",
            "confirm_password": "testing123",
        },
    )
    assert response.status_code == 201, response.text
    return {"Authorization": f'Bearer {response.json()["access_token"]}'}


def profile_payload():
    return {
        "degree": "B.Tech",
        "specialization": "Computer Science",
        "institution": "Example University",
        "graduation_year": "2026",
        "qualification": "Bachelor's",
        "technical_skills": ["Python", "SQL", "React", "Git", "JavaScript"],
        "soft_skills": ["Communication"],
        "work_experience": "Fresher (0 years)",
    }


def test_auth_profile_prediction_history_and_isolation(client):
    assert client.post("/api/intelligence/predict").status_code == 401
    headers = account(client)
    assert client.post("/api/intelligence/predict", headers=headers).status_code == 422
    assert (
        client.post(
            "/api/profile/me", headers=headers, json=profile_payload()
        ).status_code
        == 200
    )
    response = client.post("/api/intelligence/predict?k=5", headers=headers)
    assert response.status_code == 200, response.text
    result = response.json()
    assert len(result["recommendations"]) == 5
    assert len({r["career"] for r in result["recommendations"]}) == 5
    assert all(
        0 <= r["confidence"] <= 100 and 0 <= r["alignment"] <= 100
        for r in result["recommendations"]
    )
    assert all(
        set(r["matched_skills"]).isdisjoint(r["missing_skills"])
        for r in result["recommendations"]
    )
    assert (
        client.get("/api/intelligence/history", headers=headers).json()[0]["is_current"]
        is True
    )
    other = account(client, "other@example.com")
    assert client.get("/api/intelligence/history", headers=other).json() == []
    assert (
        client.post("/api/intelligence/predict?k=99", headers=headers).status_code
        == 422
    )
    client.put("/api/profile/me", headers=headers, json={"technical_skills": ["Excel"]})
    assert (
        client.get("/api/intelligence/history", headers=headers).json()[0]["is_current"]
        is False
    )


def test_profile_validation_and_login(client):
    headers = account(client)
    assert (
        client.post(
            "/api/profile/me", headers=headers, json={"graduation_year": "abc"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/profile/me", headers=headers, json={"technical_skills": [""]}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/auth/login",
            json={"email": "test@example.com", "password": "incorrect"},
        ).status_code
        == 401
    )
    assert (
        client.post(
            "/api/auth/login",
            json={"email": "test@example.com", "password": "testing123"},
        ).status_code
        == 200
    )


def test_resume_upload_review_analysis_and_ownership(client):
    from docx import Document

    headers = account(client)
    doc = Document()
    doc.add_paragraph(
        "Test Student\ntest@example.com\nSummary: Software engineer with 2 years of experience."
    )
    doc.add_paragraph(
        "Bachelor of Engineering at Example University. Python, SQL, JavaScript, React, Git."
    )
    doc.add_paragraph("Projects: built a dashboard for 100 users. github.com/example")
    buffer = io.BytesIO()
    doc.save(buffer)
    upload = client.post(
        "/api/resume/upload",
        headers=headers,
        files={
            "file": (
                "test.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 200, upload.text
    result = upload.json()
    assert "python" in result["parsed_information"]["technical_skills"]
    assert result["parsed_information"]["education"]
    assert result["parsed_information"]["parser_mode"] == "spacy_ner_and_rules"
    assert (
        client.get("/api/intelligence/resume-analysis", headers=headers).json()["score"]
        >= 75
    )
    other = account(client, "other@example.com")
    path = f'/api/resume/{result["resume_id"]}/apply'
    assert (
        client.post(
            path, headers=other, json={"technical_skills": ["Python"]}
        ).status_code
        == 404
    )
    assert (
        client.post(
            path,
            headers=headers,
            json={
                "technical_skills": ["Python", "SQL"],
                "education": ["Bachelor of Engineering"],
                "job_roles": ["Software Engineer"],
            },
        ).status_code
        == 200
    )
    assert (
        client.get("/api/profile/me", headers=headers).json()["qualification"]
        == "Bachelor of Engineering"
    )
    assert (
        client.post(
            "/api/resume/upload", headers=headers, files={"file": ("bad.exe", b"abc")}
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/resume/upload",
            headers=headers,
            files={"file": ("bad.pdf", b"not a pdf")},
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/resume/upload",
            headers=headers,
            files={"file": ("large.pdf", b"x" * (5 * 1024 * 1024 + 1))},
        ).status_code
        == 400
    )


def test_spacy_blank_fallback(monkeypatch):
    import spacy
    from app import parser

    def missing(*args, **kwargs):
        raise OSError("missing model")

    monkeypatch.setattr(spacy, "load", missing)
    parser._load_spacy.cache_clear()
    result = parser.parse_resume(
        "JS and Python. Bachelor of Science. Software Engineer. 2 years of experience."
    )
    assert result["parser_mode"] == "spacy_rules_only"
    assert {"python", "javascript"}.issubset(result["technical_skills"])
    assert result["education"] and result["job_roles"] and result["experience"]
    parser._load_spacy.cache_clear()


def test_empty_and_unknown_skills():
    with pytest.raises(ValueError):
        predict({"technical_skills": []})
    result = predict({"technical_skills": ["Unrecognized skill zzzzz"]})
    assert result["coverage"] == 0
    assert result["unrecognized_skills"]
    assert resume_analysis("", {})["score"] == 0
    assert canonical("React.js") == "react"


def test_feature_vectorization_and_report_integrity():
    import numpy as np
    from ml.features import SkillFeatures
    from app.intelligence import model_report, ROOT

    features = SkillFeatures(semantic=True).fit(["python sql", "react javascript"])
    assert np.allclose(
        features.transform(["React.js"]).toarray(),
        features.transform(["react"]).toarray(),
    )
    assert features.transform(["python"]).shape[1] > 384
    report = model_report()
    split = json.loads((ROOT / "artifacts/split.json").read_text())
    assert set(split["train"]).isdisjoint(split["test"])
    assert len(split["train"]) == report["train_rows"]
    assert len(split["test"]) == report["test_rows"]
    assert report["embedding"]["training_pairs"] == 480
